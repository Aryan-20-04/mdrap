"""
Unit and Integration Tests for MDRAP Excel (.xlsx) and CSV Financial Model Exporter.
"""
import csv
import os
import shutil
import sys
import tempfile
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from exporter import MarketDataExporter
from storage import Store
from pipeline import Pipeline
from simulator import FeedSimulator, SimulatorConfig


@pytest.fixture
def temp_env():
    """Create isolated temporary directory for test outputs and databases."""
    temp_dir = tempfile.mkdtemp(prefix="mdrap_export_test_")
    db_path = os.path.join(temp_dir, "test_mdrap.db")
    
    # Pre-populate database with a controlled run
    store = Store(db_path)
    pipe = Pipeline(store=store)
    sim = FeedSimulator(SimulatorConfig(seed=42, num_events=1000))
    for raw, _label in sim.generate():
        pipe.process_one(raw)
    pipe.finish()
    store.close()

    yield {
        "dir": temp_dir,
        "db": db_path,
    }

    try:
        shutil.rmtree(temp_dir, ignore_errors=True)
    except Exception:
        pass


def test_exporter_gather_data(temp_env):
    """Verify gather_report_data returns structured market microstructure metrics."""
    exporter = MarketDataExporter(db_path=temp_env["db"])
    data = exporter.gather_report_data("AAPL")

    assert data["symbol"] == "AAPL"
    assert "timestamp" in data
    assert "bbo" in data
    assert "ladder" in data
    assert "vwap_curve" in data
    assert "candles" in data
    assert "health" in data
    assert data["total_canonical"] > 0


def test_exporter_excel_generation(temp_env):
    """Verify export_excel creates a valid 5-sheet workbook with expected cell content."""
    import openpyxl

    exporter = MarketDataExporter(db_path=temp_env["db"])
    out_xlsx = os.path.join(temp_env["dir"], "test_report.xlsx")
    res_path = exporter.export_excel("AAPL", output_path=out_xlsx)

    assert os.path.exists(res_path)
    assert os.path.getsize(res_path) > 1000

    wb = openpyxl.load_workbook(res_path)
    expected_sheets = [
        "Executive Overview",
        "VWAP Slippage Model",
        "Consolidated L2 Depth",
        "OHLCV Market Candles",
        "Data Quality & Audit",
    ]
    assert wb.sheetnames == expected_sheets

    # Tab 1 checks
    ws1 = wb["Executive Overview"]
    assert "MDRAP Institutional Market Summary: AAPL" in str(ws1["A1"].value)
    assert ws1.max_row >= 10
    assert ws1.max_column >= 2

    # Tab 2 checks (VWAP)
    ws2 = wb["VWAP Slippage Model"]
    assert "Institutional VWAP" in str(ws2["A1"].value)
    assert ws2.max_row >= 10

    # Tab 3 checks (L2 Depth)
    ws3 = wb["Consolidated L2 Depth"]
    assert "Consolidated Multi-Venue Order Book Depth" in str(ws3["A1"].value)
    assert ws3.max_column == 9

    # Tab 4 checks (OHLCV)
    ws4 = wb["OHLCV Market Candles"]
    assert "OHLCV" in str(ws4["A1"].value)

    # Tab 5 checks (Quality & SLA Audit)
    ws5 = wb["Data Quality & Audit"]
    assert "Quality & Feed SLA Audit" in str(ws5["A1"].value)
    assert ws5.max_row >= 5
    wb.close()


def test_exporter_csv_generation(temp_env):
    """Verify export_csv creates all 6 standalone structured CSV files."""
    exporter = MarketDataExporter(db_path=temp_env["db"])
    csv_dir = os.path.join(temp_env["dir"], "csv_out")
    files = exporter.export_csv("AAPL", output_dir=csv_dir)

    assert len(files) == 6
    basenames = [os.path.basename(f) for f in files]
    assert "overview.csv" in basenames
    assert "vwap_slippage.csv" in basenames
    assert "l2_depth.csv" in basenames
    assert "ohlcv_candles.csv" in basenames
    assert "feed_health.csv" in basenames
    assert "quarantine_sample.csv" in basenames

    # Check overview.csv content
    with open(os.path.join(csv_dir, "overview.csv"), "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
        assert len(rows) >= 5
        assert rows[0] == ["metric", "value"]

    # Check vwap_slippage.csv headers
    with open(os.path.join(csv_dir, "vwap_slippage.csv"), "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        assert "target_size" in header
        assert "vwap_price" in header
        assert "slippage_dollars" in header


def test_exporter_auto_populates_empty_db(temp_env):
    """Verify that an empty DB is automatically populated via simulation if queried."""
    empty_db = os.path.join(temp_env["dir"], "empty.db")
    exporter = MarketDataExporter(db_path=empty_db)
    
    out_xlsx = os.path.join(temp_env["dir"], "empty_auto.xlsx")
    res_path = exporter.export_excel("AAPL", output_path=out_xlsx)
    assert os.path.exists(res_path)
    assert os.path.getsize(res_path) > 1000

    # Verify DB now contains events
    store = Store(empty_db)
    cnts = store.counts()
    assert sum(cnts.values()) >= 500
    store.close()


def test_cli_export_command(temp_env):
    """Verify CLI parser dispatches 'export' subcommand correctly."""
    from cli import build_parser

    parser = build_parser()
    out_file = os.path.join(temp_env["dir"], "cli_out.xlsx")
    args = parser.parse_args(["export", "AAPL", "--db", temp_env["db"], "-o", out_file])
    assert args.func is not None
    assert args.symbol == "AAPL"
    assert args.db == temp_env["db"]
    assert args.output == out_file

    # Execute func directly
    args.func(args)
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000


def test_export_data_json_and_sql_injection_defense(temp_env):
    """Verify export_data exports valid JSON and defends against SQL injection."""
    from export import export_data

    # 1. Valid export to JSON
    json_out = os.path.join(temp_env["dir"], "valid.json")
    count = export_data(temp_env["db"], table="canonical_events", output_path=json_out, fmt="json")
    assert count > 0
    assert os.path.exists(json_out)

    # 2. SQL injection attempt in table name rejected
    malicious_out = os.path.join(temp_env["dir"], "hacked.json")
    with pytest.raises(ValueError, match="does not exist in database"):
        export_data(
            temp_env["db"],
            table="canonical_events; DROP TABLE canonical_events; --",
            output_path=malicious_out,
            fmt="json",
        )
