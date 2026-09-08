import os
import sys
import tempfile
import time
import urllib.request
import pytest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from excel_bridge import (
    GLOBAL_MARKET_STATE,
    ExcelBridgeServer,
    generate_bloomberg_replacement_workbook,
)


def test_market_data_state_get_field():
    state = GLOBAL_MARKET_STATE
    # Check default tickers
    px = state.get_field("AAPL", "PX_LAST")
    assert float(px) > 0

    bid = state.get_field("AAPL", "BID")
    ask = state.get_field("AAPL", "ASK")
    assert float(bid) < float(ask)

    # Vendor suffix stripping
    px_suffix = state.get_field("AAPL US Equity", "PX_LAST")
    assert px_suffix == px

    # CVD & Flow fields
    cvd = state.get_field("AAPL", "CVD")
    assert cvd != "#N/A"

    buy_pct = state.get_field("AAPL", "BUY_PCT")
    assert "%" in buy_pct

    stance = state.get_field("AAPL", "INST_STANCE")
    assert len(stance) > 0

    # TCA score
    tca_score = state.get_field("AAPL", "BEST_EX_SCORE")
    assert float(tca_score) > 0

    # Invalid ticker / field
    err_sym = state.get_field("NONEXISTENT", "PX_LAST")
    assert "#N/A Symbol" in err_sym

    err_fld = state.get_field("AAPL", "INVALID_FIELD_XYZ")
    assert "#N/A Field" in err_fld


def test_excel_bridge_server_endpoints():
    port = 18095
    server = ExcelBridgeServer(host="127.0.0.1", port=port)
    server.start(daemon=True)
    time.sleep(0.3)

    try:
        base = f"http://127.0.0.1:{port}"

        # 1. Test /bdp
        with urllib.request.urlopen(f"{base}/bdp?ticker=AAPL&field=PX_LAST", timeout=3) as resp:
            assert resp.status == 200
            val = resp.read().decode("utf-8")
            assert float(val) > 0

        # 2. Test /bdp CVD
        with urllib.request.urlopen(f"{base}/bdp?ticker=AAPL&field=CVD", timeout=3) as resp:
            assert resp.status == 200
            val = resp.read().decode("utf-8")
            assert int(val.replace("+", "")) != 0 or val == "0"

        # 3. Test /api/quote
        with urllib.request.urlopen(f"{base}/api/quote?symbol=AAPL", timeout=3) as resp:
            assert resp.status == 200
            import json
            data = json.loads(resp.read().decode("utf-8"))
            assert "quote" in data
            assert data["quote"]["symbol"] == "AAPL"

        # 4. Test /live.csv
        with urllib.request.urlopen(f"{base}/live.csv", timeout=3) as resp:
            assert resp.status == 200
            content = resp.read().decode("utf-8")
            assert "Symbol,LastPrice,Bid,Ask" in content
            assert "AAPL" in content

        # 5. Test /vba
        with urllib.request.urlopen(f"{base}/vba", timeout=3) as resp:
            assert resp.status == 200
            vba = resp.read().decode("utf-8")
            assert "Public Function BDP" in vba

    finally:
        server.stop()


def test_generate_bloomberg_replacement_workbook():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        out = generate_bloomberg_replacement_workbook(output_path=path, port=8085)
        assert os.path.exists(out)
        wb = openpyxl.load_workbook(out, data_only=False)
        assert "Live Market Desk" in wb.sheetnames
        assert "Broker Flow & MPID Matrix" in wb.sheetnames
        assert "Bloomberg BDP Cheat Sheet" in wb.sheetnames

        ws1 = wb["Live Market Desk"]
        # Formula check
        assert "=NUMBERVALUE(WEBSERVICE(" in str(ws1["C5"].value)
    finally:
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
